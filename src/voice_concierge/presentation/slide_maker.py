# src/voice_concierge/presentation/slide_maker.py
from typing import Dict, Any, List, Optional
import os
from pydantic import BaseModel, Field
from tenacity import retry, stop_after_attempt, wait_exponential
from pptx import Presentation

from agents import Tool
from voice_concierge.utils.logging_utils import get_logger

logger = get_logger(__name__)


class BulletPoint(BaseModel):
    """Model for a bullet point in a slide."""

    text: str = Field(..., description="Text content of the bullet point")
    level: int = Field(0, description="Indentation level (0 for main points)")


class SlideContent(BaseModel):
    """Model for slide content with title and bullet points."""

    title: str = Field(..., description="Slide title")
    bullets: List[BulletPoint] = Field(
        default_factory=list, description="Bullet points"
    )


class PresentationContent(BaseModel):
    """Model for overall presentation content."""

    title: str = Field(..., description="Presentation title")
    slides: List[SlideContent] = Field(..., description="List of slides")


class SlideMaker(Tool):
    """Tool for creating PowerPoint presentations from structured content."""

    name = "make_slides"
    description = "Create PowerPoint presentations from bullet point content"

    def __init__(self, template_path: Optional[str] = None):
        """Initialize the slide maker tool.

        Args:
            template_path: Optional path to a PowerPoint template to use.
        """
        self.template_path = template_path
        logger.info(f"Initialized SlideMaker with template: {template_path or 'None'}")

    def _create_title_slide(self, prs: Presentation, title: str) -> None:
        """Create the title slide for the presentation.

        Args:
            prs: Presentation object
            title: Presentation title
        """
        slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title slide layout
        title_shape = slide.shapes.title
        subtitle_shape = slide.placeholders[1]

        title_shape.text = title
        subtitle_shape.text = f"Generated on {datetime.now().strftime('%B %d, %Y')}"

    def _create_content_slide(self, prs: Presentation, content: SlideContent) -> None:
        """Create a content slide with title and bullet points.

        Args:
            prs: Presentation object
            content: Slide content with title and bullets
        """
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content layout

        # Set the slide title
        title_shape = slide.shapes.title
        title_shape.text = content.title

        # Add bullet points
        body_shape = slide.placeholders[1]
        text_frame = body_shape.text_frame

        # Clear any default text
        if text_frame.paragraphs:
            p = text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = ""

        # Add the bullet points
        for i, bullet in enumerate(content.bullets):
            if i == 0 and text_frame.paragraphs and not text_frame.paragraphs[0].runs:
                # Use the first paragraph if it's empty
                p = text_frame.paragraphs[0]
            else:
                # Create a new paragraph
                p = text_frame.add_paragraph()

            # Set indentation level
            p.level = bullet.level

            # Set text
            p.text = bullet.text

    @retry(
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=2, max=10),
        reraise=True,
    )
    def run(self, content: Dict[str, Any], output_path: str = "summary.pptx") -> str:
        """Create a PowerPoint presentation from structured content.

        Args:
            content: Structured content for the presentation (can be dict or JSON string)
            output_path: Path where the PowerPoint file should be saved

        Returns:
            Absolute path to the created PowerPoint file
        """
        logger.info(f"Creating presentation at: {output_path}")

        try:
            # Parse the input content
            if isinstance(content, str):
                content = json.loads(content)

            presentation_content = PresentationContent(**content)

            # Create the presentation
            if self.template_path and os.path.exists(self.template_path):
                prs = Presentation(self.template_path)
                logger.debug(f"Using template: {self.template_path}")
            else:
                prs = Presentation()
                logger.debug("Using default blank presentation")

            # Create title slide
            self._create_title_slide(prs, presentation_content.title)

            # Create content slides
            for slide_content in presentation_content.slides:
                self._create_content_slide(prs, slide_content)

            # Save the presentation
            prs.save(output_path)

            # Return the absolute path
            abs_path = os.path.abspath(output_path)
            logger.info(f"Presentation created successfully: {abs_path}")
            return abs_path

        except Exception as e:
            logger.error(f"Error creating presentation: {str(e)}")
            raise
